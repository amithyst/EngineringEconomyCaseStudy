#!/usr/bin/env python3

"""Export an Excel workbook into a single Markdown document."""

from __future__ import annotations

import argparse
import os
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_DIR = REPO_ROOT / "outputs" / "markdown"


@dataclass(frozen=True)
class Bounds:
    min_row: int
    max_row: int
    min_col: int
    max_col: int


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Read all worksheet data from an .xlsx/.xlsm file and export it to one Markdown file."
    )
    parser.add_argument("input_path", type=Path, help="Workbook path (.xlsx/.xlsm)")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Output Markdown path. Defaults to outputs/markdown/<workbook-name>.md",
    )
    return parser.parse_args()


def resolve_path(path: Path) -> Path:
    candidate = path if path.is_absolute() else (Path.cwd() / path)
    return candidate.resolve()


def detect_bounds(worksheet) -> Bounds | None:
    coords: list[tuple[int, int]] = []
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                coords.append((cell.row, cell.column))

    if not coords:
        return None

    rows = [row for row, _ in coords]
    cols = [col for _, col in coords]
    return Bounds(min(rows), max(rows), min(cols), max(cols))


def normalize_scalar(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value)


def excel_date_to_text(value) -> str:
    if isinstance(value, datetime):
        if value.time() == time(0, 0, 0):
            return f"{value.year}/{value.month}/{value.day}"
        return f"{value.year}/{value.month}/{value.day} {value.hour}:{value.minute:02d}:{value.second:02d}"
    if isinstance(value, date):
        return f"{value.year}/{value.month}/{value.day}"
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return normalize_scalar(value)


def format_number(value, number_format: str) -> str:
    if value is None:
        return ""
    if not isinstance(value, (int, float, Decimal)) or isinstance(value, bool):
        return normalize_scalar(value)

    number_format = (number_format or "General").strip()
    if number_format == "General":
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)

    if "%" in number_format:
        decimals = 0
        match = re.search(r"\.([0#]+)%", number_format)
        if match:
            decimals = len(match.group(1))
        return f"{float(value) * 100:.{decimals}f}%"

    comma = "," if "," in number_format else ""
    decimals = 0
    match = re.search(r"\.([0#]+)", number_format)
    if match:
        decimals = len(match.group(1))

    if decimals == 0:
        return f"{float(value):{comma}.0f}"
    return f"{float(value):{comma}.{decimals}f}"


def build_merged_anchor_map(worksheet) -> dict[tuple[int, int], tuple[int, int]]:
    anchor_map: dict[tuple[int, int], tuple[int, int]] = {}
    for merged in worksheet.merged_cells.ranges:
        cell_range = CellRange(str(merged))
        anchor = (cell_range.min_row, cell_range.min_col)
        for row in range(cell_range.min_row, cell_range.max_row + 1):
            for col in range(cell_range.min_col, cell_range.max_col + 1):
                anchor_map[(row, col)] = anchor
    return anchor_map


def format_cell(
    formula_cell: Cell,
    value_cell: Cell,
    merged_anchor_map: dict[tuple[int, int], tuple[int, int]],
    formula_ws,
    value_ws,
) -> str:
    anchor = merged_anchor_map.get((formula_cell.row, formula_cell.column))
    if anchor and anchor != (formula_cell.row, formula_cell.column):
        formula_cell = formula_ws.cell(*anchor)
        value_cell = value_ws.cell(*anchor)

    raw_display = value_cell.value
    if is_date_format(formula_cell.number_format):
        display_value = excel_date_to_text(raw_display)
    else:
        display_value = format_number(raw_display, formula_cell.number_format)

    raw_value = formula_cell.value

    if isinstance(raw_value, str) and raw_value.startswith("="):
        if display_value:
            return f"{display_value} `({raw_value})`"
        return f"`{raw_value}`"

    return display_value


def escape_md(text: str) -> str:
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def build_table(headers: list[str], rows: Iterable[list[str]]) -> str:
    lines = [
        "| " + " | ".join(escape_md(header) for header in headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(escape_md(cell) for cell in row) + " |")
    return "\n".join(lines)


def sheet_to_markdown(formula_ws, value_ws) -> str:
    bounds = detect_bounds(formula_ws)
    lines = [f"## Sheet: {formula_ws.title}", ""]

    if bounds is None:
        lines.append("This sheet is empty.")
        return "\n".join(lines)

    merged_ranges = [str(item) for item in formula_ws.merged_cells.ranges]
    lines.extend(
        [
            f"- Range: `{get_column_letter(bounds.min_col)}{bounds.min_row}:{get_column_letter(bounds.max_col)}{bounds.max_row}`",
            f"- Rows: `{bounds.max_row - bounds.min_row + 1}`",
            f"- Columns: `{bounds.max_col - bounds.min_col + 1}`",
            f"- Merged cells: `{', '.join(merged_ranges)}`" if merged_ranges else "- Merged cells: none",
            "",
        ]
    )

    headers = ["Row"] + [get_column_letter(col) for col in range(bounds.min_col, bounds.max_col + 1)]
    merged_anchor_map = build_merged_anchor_map(formula_ws)
    table_rows: list[list[str]] = []
    for row_idx in range(bounds.min_row, bounds.max_row + 1):
        row = [str(row_idx)]
        for col_idx in range(bounds.min_col, bounds.max_col + 1):
            row.append(
                format_cell(
                    formula_ws.cell(row=row_idx, column=col_idx),
                    value_ws.cell(row=row_idx, column=col_idx),
                    merged_anchor_map,
                    formula_ws,
                    value_ws,
                )
            )
        table_rows.append(row)

    lines.append(build_table(headers, table_rows))
    lines.append("")
    return "\n".join(lines)


def workbook_to_markdown(input_path: Path) -> str:
    formula_wb = load_workbook(input_path, data_only=False, keep_vba=True)
    value_wb = load_workbook(input_path, data_only=True, keep_vba=True)

    lines = [
        f"# {input_path.name}",
        "",
        f"- Source file: `{input_path}`",
        f"- Sheet count: `{len(formula_wb.worksheets)}`",
        f"- Exported at: `{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}`",
        "",
    ]

    for formula_ws, value_ws in zip(formula_wb.worksheets, value_wb.worksheets):
        lines.append(sheet_to_markdown(formula_ws, value_ws))

    return "\n".join(lines).strip() + "\n"


def default_output_path(input_path: Path) -> Path:
    return DEFAULT_OUTPUT_DIR / f"{input_path.stem}.md"


def main() -> None:
    args = parse_args()
    input_path = resolve_path(args.input_path)

    if not input_path.exists():
        raise FileNotFoundError(f"Workbook not found: {input_path}")

    output_path = resolve_path(args.output) if args.output else default_output_path(input_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    markdown = workbook_to_markdown(input_path)
    output_path.write_text(markdown, encoding="utf-8")

    rel_input = os.path.relpath(input_path, Path.cwd())
    rel_output = os.path.relpath(output_path, Path.cwd())
    print(f"Converted {rel_input} -> {rel_output}")


if __name__ == "__main__":
    main()
