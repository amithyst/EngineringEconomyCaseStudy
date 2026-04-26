#!/usr/bin/env python3

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


REPO_ROOT = Path(__file__).resolve().parents[2]
WORKBOOK_PATH = next(
    path for path in (REPO_ROOT / "final_report").glob("*.xlsx")
    if "投资分析" in path.name
)


def q(sheet_name: str) -> str:
    return f"'{sheet_name}'"


def set_formula(cell, formula: str, number_format: str | None = None) -> None:
    cell.value = formula
    if number_format:
        cell.number_format = number_format


def apply_basic_table_style(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def formulaize_workbook() -> None:
    wb = load_workbook(WORKBOOK_PATH)
    cover, ass, annual, quarterly, valuation, sensitivity = wb.worksheets[:6]
    ass_name, annual_name, helper_name = ass.title, annual.title, "敏感性计算明细"

    # Cover: reflect the added calculation-detail sheet.
    cover["B13"] = '  附录：敏感性计算明细（见"敏感性计算明细"工作表）'
    cover["B13"].font = Font(name="Microsoft YaHei", size=10, color="1F4E79")

    # Assumptions: totals and key percentage input as auditable inputs.
    for col in "CDEFG":
        set_formula(ass[f"{col}13"], f"=SUM({col}5:{col}12)", "#,##0")
    ass["C43"] = 0.025
    ass["C43"].number_format = "0.0%"

    # Annual cash flow formulas.
    year_cols = ["C", "D", "E", "F"]
    assumption_rows = [27, 28, 29, 30]
    parking_rows = [34, 35, 36, 37]
    cost_growth_cols = ["C", "D", "E", "F"]
    for idx, col in enumerate(year_cols):
        arow = assumption_rows[idx]
        prow = parking_rows[idx]
        months_formula = f"(12-{q(ass_name)}!$C$64)" if idx == 0 else "12"
        rent_growth_formula = f"(1+{q(ass_name)}!D{arow})" if idx < 3 else f"(1+{q(ass_name)}!D{arow})"
        set_formula(
            annual[f"{col}6"],
            f"=SUMPRODUCT({q(ass_name)}!$E$5:$E$12,{q(ass_name)}!$C$16:$C$23)*"
            f"{q(ass_name)}!C{arow}*{q(ass_name)}!E{arow}*{rent_growth_formula}/10000",
            "#,##0.00",
        )
        set_formula(
            annual[f"{col}7"],
            f"=({q(ass_name)}!$F$13+{q(ass_name)}!$G$13)*{q(ass_name)}!C{prow}*"
            f"{q(ass_name)}!D{prow}*{q(ass_name)}!E{prow}*{q(ass_name)}!F{arow}/10000",
            "#,##0.00",
        )
        set_formula(
            annual[f"{col}8"],
            f"={q(ass_name)}!$E$13*{q(ass_name)}!C{arow}*10*{months_formula}/10000",
            "#,##0.00",
        )
        set_formula(annual[f"{col}9"], f"=SUM({col}6:{col}8)", "#,##0.00")
        set_formula(annual[f"{col}14"], f"={col}6*{q(ass_name)}!$C$43", "#,##0.00")
        set_formula(annual[f"{col}18"], f"=SUM({col}11:{col}17)", "#,##0.00")
        set_formula(annual[f"{col}20"], f"={col}9-{col}18", "#,##0.00")
        set_formula(annual[f"{col}21"], f"=MAX({col}20,0)*{q(ass_name)}!$C$61", "#,##0.00")
        set_formula(annual[f"{col}22"], f"={col}20-{col}21", "#,##0.00")
        set_formula(annual[f"{col}26"], f"={col}22", "#,##0.00")
        set_formula(annual[f"{col}29"], f"={q(ass_name)}!C{arow}", "0.0%")
        set_formula(annual[f"{col}30"], f"={q(ass_name)}!F{arow}", "0")
        set_formula(annual[f"{col}31"], f"={col}6/{col}9", "0.0%")
        set_formula(annual[f"{col}32"], f"={col}20/{col}9", "0.0%")

    # Fixed operating cost schedule.
    for row, input_row, growth_row, uses_area in [
        (11, 40, 50, False),
        (12, 41, 51, False),
        (13, 42, 52, True),
        (15, 44, 53, True),
        (16, 45, 54, True),
        (17, 46, 55, False),
    ]:
        base = f"{q(ass_name)}!$C$13*{q(ass_name)}!$C${input_row}/10000" if uses_area else f"{q(ass_name)}!$C${input_row}/10000"
        set_formula(annual[f"C{row}"], f"={base}*(12-{q(ass_name)}!$C$64)/12", "#,##0.00")
        set_formula(annual[f"D{row}"], f"={base}*(1+{q(ass_name)}!D{growth_row})", "#,##0.00")
        set_formula(annual[f"E{row}"], f"=D{row}*(1+{q(ass_name)}!E{growth_row})", "#,##0.00")
        set_formula(annual[f"F{row}"], f"=E{row}*(1+{q(ass_name)}!F{growth_row})", "#,##0.00")

    set_formula(annual["F24"], f"=F22/{q(ass_name)}!$C$63", "#,##0.00")
    set_formula(annual["F26"], "=F22+F24", "#,##0.00")

    # Quarterly cash flow formulas.
    quarter_days = {"C": 0, "D": 91, "E": 92, "F": 92}
    for col, days in quarter_days.items():
        if col in ("C", "D"):
            quarterly[f"{col}6"] = 0
        else:
            set_formula(quarterly[f"{col}6"], f"={q(annual_name)}!$C$6/{q(ass_name)}!$E$27*{days}", "#,##0.00")
        set_formula(
            quarterly[f"{col}7"],
            f"=({q(ass_name)}!$F$13+{q(ass_name)}!$G$13)*{q(ass_name)}!$C$34*"
            f"{q(ass_name)}!$D$34*{q(ass_name)}!$E$34*{days}/10000",
            "#,##0.00",
        )
        if col == "C":
            quarterly[f"{col}8"] = 0
            quarterly[f"{col}11"] = 0
        else:
            set_formula(quarterly[f"{col}8"], f"={q(annual_name)}!$C$8/3", "#,##0.00")
            set_formula(quarterly[f"{col}11"], f"={q(annual_name)}!$C$18/3", "#,##0.00")
        set_formula(quarterly[f"{col}9"], f"=SUM({col}6:{col}8)", "#,##0.00")
        set_formula(quarterly[f"{col}13"], f"={col}9-{col}11", "#,##0.00")
    set_formula(quarterly["F16"], f"={q(annual_name)}!$C$22", "#,##0.00")

    # Valuation formulas.
    for row, col, exponent in zip(range(5, 9), year_cols, range(1, 5)):
        set_formula(valuation[f"D{row}"], f"={q(annual_name)}!{col}26", "#,##0.00")
        set_formula(valuation[f"E{row}"], f"=1/(1+{q(ass_name)}!$C$62)^{exponent}", "0.0000")
        set_formula(valuation[f"F{row}"], f"=D{row}*E{row}", "#,##0.00")
    set_formula(valuation["F9"], "=SUM(F5:F8)", "#,##0.00")
    set_formula(valuation["C12"], "=F9", "#,##0.00")
    set_formula(valuation["C13"], f"=-{q(ass_name)}!$C$59/10000", "#,##0.00")
    set_formula(valuation["C14"], f"=-{q(ass_name)}!$C$60/10000", "#,##0.00")
    set_formula(valuation["C15"], "=SUM(C12:C14)", "#,##0.00")
    set_formula(valuation["C16"], f"=1+{q(ass_name)}!$C$58", "0.00")
    set_formula(valuation["C17"], "=C15/C16", "#,##0.00")
    set_formula(valuation["C18"], f"=C17*{q(ass_name)}!$C$58", "#,##0.00")
    set_formula(valuation["C19"], f"={q(ass_name)}!$C$59/10000", "#,##0.00")
    set_formula(valuation["C20"], f"={q(ass_name)}!$C$60/10000", "#,##0.00")
    set_formula(valuation["C21"], f"=C17*(1+{q(ass_name)}!$C$58)+C19+C20", "#,##0.00")
    set_formula(valuation["C24"], "=F9-C21", "#,##0.00")
    set_formula(valuation["C25"], f"={q(ass_name)}!$C$62", "0%")
    set_formula(valuation["C26"], "=F9/C21", "0.00")
    set_formula(valuation["C27"], f"={q(annual_name)}!F24/C21", "0.00")
    set_formula(valuation["C28"], "=F9/C21", "0.00")

    # Sensitivity detail sheet.
    if helper_name in wb.sheetnames:
        del wb[helper_name]
    helper = wb.create_sheet(helper_name, 6)
    helper.sheet_view.showGridLines = False
    helper.freeze_panes = "A4"
    helper["A1"] = "敏感性计算明细：稳定期出租率 × 2029年租金增长率"
    helper["A1"].font = Font(name="Microsoft YaHei", size=13, bold=True, color="FFFFFF")
    helper["A1"].fill = PatternFill("solid", fgColor="1F3A5F")
    helper.merge_cells("A1:AM1")
    headers = [
        "情景", "稳定期出租率", "2029租金增长率", "2026出租率",
        "租金2026", "租金2027", "租金2028", "租金2029",
        "停车2026", "停车2027", "停车2028", "停车2029",
        "物业2026", "物业2027", "物业2028", "物业2029",
        "收入2026", "收入2027", "收入2028", "收入2029",
        "O&M2026", "O&M2027", "O&M2028", "O&M2029",
        "BTCF2026", "BTCF2027", "BTCF2028", "BTCF2029",
        "Tax2026", "Tax2027", "Tax2028", "Tax2029",
        "ATCF2026", "ATCF2027", "ATCF2028", "ATCF2029",
        "退出价值", "未来CF现值", "建议收购对价",
    ]
    for c, header in enumerate(headers, 1):
        cell = helper.cell(3, c, header)
        cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E75B6")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    occupancies = [0.70, 0.75, 0.80, 0.85, 0.90, 0.95]
    growths = [-0.03, 0.00, 0.03, 0.05, 0.08]
    row = 4
    for occ in occupancies:
        for growth in growths:
            helper.cell(row, 1, f"S{row-3:02d}")
            helper.cell(row, 2, occ).number_format = "0%"
            helper.cell(row, 3, growth).number_format = "0%"
            set_formula(helper.cell(row, 4), f"=B{row}-15%", "0%")
            occ_refs = [f"D{row}", f"B{row}", f"B{row}", f"B{row}"]
            growth_refs = ["0", "0", "0", f"C{row}"]
            day_rows = [27, 28, 29, 30]
            # Rent E:H
            for offset, col_idx in enumerate(range(5, 9)):
                set_formula(
                    helper.cell(row, col_idx),
                    f"=SUMPRODUCT({q(ass_name)}!$E$5:$E$12,{q(ass_name)}!$C$16:$C$23)*"
                    f"{occ_refs[offset]}*{q(ass_name)}!E{day_rows[offset]}*(1+{growth_refs[offset]})/10000",
                    "#,##0.00",
                )
            # Parking I:L
            for offset, col_idx in enumerate(range(9, 13)):
                prow = 34 + offset
                set_formula(
                    helper.cell(row, col_idx),
                    f"=({q(ass_name)}!$F$13+{q(ass_name)}!$G$13)*{q(ass_name)}!C{prow}*"
                    f"{q(ass_name)}!D{prow}*{q(ass_name)}!E{prow}*{q(ass_name)}!F{day_rows[offset]}/10000",
                    "#,##0.00",
                )
            # Property M:P
            for offset, col_idx in enumerate(range(13, 17)):
                months = f"(12-{q(ass_name)}!$C$64)" if offset == 0 else "12"
                set_formula(
                    helper.cell(row, col_idx),
                    f"={q(ass_name)}!$E$13*{occ_refs[offset]}*10*{months}/10000",
                    "#,##0.00",
                )
            # Income Q:T
            for offset, col_idx in enumerate(range(17, 21)):
                rent_col = chr(ord("E") + offset)
                park_col = chr(ord("I") + offset)
                prop_col = chr(ord("M") + offset)
                set_formula(helper.cell(row, col_idx), f"={rent_col}{row}+{park_col}{row}+{prop_col}{row}", "#,##0.00")
            fixed_cols = ["C", "D", "E", "F"]
            rent_cols = ["E", "F", "G", "H"]
            for offset, col_idx in enumerate(range(21, 25)):
                ac = fixed_cols[offset]
                rc = rent_cols[offset]
                set_formula(
                    helper.cell(row, col_idx),
                    f"={q(annual_name)}!{ac}$11+{q(annual_name)}!{ac}$12+{q(annual_name)}!{ac}$13+"
                    f"{q(annual_name)}!{ac}$15+{q(annual_name)}!{ac}$16+{q(annual_name)}!{ac}$17+"
                    f"{rc}{row}*{q(ass_name)}!$C$43",
                    "#,##0.00",
                )
            # BTCF, tax, ATCF
            for offset, col_idx in enumerate(range(25, 29)):
                inc_col = chr(ord("Q") + offset)
                opex_col = chr(ord("U") + offset)
                set_formula(helper.cell(row, col_idx), f"={inc_col}{row}-{opex_col}{row}", "#,##0.00")
            btc_cols = ["Y", "Z", "AA", "AB"]
            for offset, col_idx in enumerate(range(29, 33)):
                btc_col = btc_cols[offset]
                set_formula(helper.cell(row, col_idx), f"=MAX({btc_col}{row},0)*{q(ass_name)}!$C$61", "#,##0.00")
            for offset, col_idx in enumerate(range(33, 37)):
                btc_col = btc_cols[offset]
                tax_col = ["AC", "AD", "AE", "AF"][offset]
                set_formula(helper.cell(row, col_idx), f"={btc_col}{row}-{tax_col}{row}", "#,##0.00")
            set_formula(helper.cell(row, 37), f"=AJ{row}/{q(ass_name)}!$C$63", "#,##0.00")
            set_formula(
                helper.cell(row, 38),
                f"=AG{row}/(1+{q(ass_name)}!$C$62)^1+AH{row}/(1+{q(ass_name)}!$C$62)^2+"
                f"AI{row}/(1+{q(ass_name)}!$C$62)^3+(AJ{row}+AK{row})/(1+{q(ass_name)}!$C$62)^4",
                "#,##0.00",
            )
            set_formula(
                helper.cell(row, 39),
                f"=(AL{row}-{q(ass_name)}!$C$59/10000-{q(ass_name)}!$C$60/10000)/(1+{q(ass_name)}!$C$58)",
                "#,##0.00",
            )
            row += 1
    apply_basic_table_style(helper, 3, row - 1, 1, 39)
    for col in range(1, 40):
        helper.column_dimensions[helper.cell(3, col).column_letter].width = 13
    helper.column_dimensions["A"].width = 8
    helper.column_dimensions["B"].width = 13
    helper.column_dimensions["C"].width = 15
    helper.column_dimensions["AM"].width = 16

    # Sensitivity sheet formulas linking to detail sheet.
    for i, occ in enumerate(occupancies, start=5):
        sensitivity[f"B{i}"] = occ
        sensitivity[f"B{i}"].number_format = "0%"
    for j, growth in enumerate(growths, start=3):
        sensitivity.cell(4, j, growth)
        sensitivity.cell(4, j).number_format = "0%"
    for r in range(5, 11):
        for c in range(3, 8):
            set_formula(
                sensitivity.cell(r, c),
                f"=SUMIFS({q(helper_name)}!$AM:$AM,{q(helper_name)}!$B:$B,$B{r},{q(helper_name)}!$C:$C,{sensitivity.cell(4, c).coordinate})",
                "#,##0.00",
            )
    for r in range(14, 20):
        source_row = r - 9
        set_formula(sensitivity[f"C{r}"], f"=E{source_row}", "#,##0.00")
        set_formula(sensitivity[f"D{r}"], f"=(C{r}-$E$10)/$E$10", "0.0%")
    for r, c_source in zip(range(23, 28), ["C", "D", "E", "F", "G"]):
        set_formula(sensitivity[f"C{r}"], f"={c_source}10", "#,##0.00")
        set_formula(sensitivity[f"D{r}"], f"=(C{r}-$E$10)/$E$10", "0.0%")
    sensitivity["B31"] = (
        '="基准假设（出租率95%，租金增长率3%）下建议收购价约 "&TEXT($E$10,"0")&'
        '" 万元（3.33亿元）。按“2026年出租率=稳定期出租率-15个百分点”的口径，出租率每下降5个百分点，'
        '收购对价约下降"&TEXT($E$10-$E$9,"0")&"万元；若稳定出租率仅70%且2029年租金增长率为-3%，合理价格约"&TEXT($C$5,"0")&"万元。"'
    )
    sensitivity["B31"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Calculation settings.
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(WORKBOOK_PATH)


if __name__ == "__main__":
    formulaize_workbook()
